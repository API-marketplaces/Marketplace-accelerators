import csv
import os
import uuid
from dataclasses import dataclass
from fastapi import APIRouter, Depends, Form, HTTPException, status, UploadFile, File, Query
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool
from typing import Iterable, List
from openpyxl import load_workbook
# Local imports from the application's modules
from app.database.database import get_db
from app.schemas.audience import AudienceResponse
from app.schemas.file import FilePublic
from app.crud import files as crud_files
from app.core.security import get_current_user
from app.core.config import settings

# Initialize the API router for file-related endpoints
router = APIRouter()

UPLOAD_CHUNK_SIZE = 1024 * 1024


@dataclass
class FileInspection:
    columns: List[str]
    records_count: int | None


def _normalize_columns(columns: Iterable[object]) -> List[str]:
    return [str(col).strip() for col in columns if col is not None and str(col).strip()]


def _max_upload_bytes() -> int:
    return settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024


def _is_safe_upload_path(file_path: str) -> bool:
    upload_dir_abs = os.path.abspath(settings.UPLOAD_DIRECTORY)
    candidate = os.path.abspath(file_path)
    try:
        return os.path.commonpath([upload_dir_abs, candidate]) == upload_dir_abs
    except ValueError:
        return False


def _copy_upload_to_disk(file: UploadFile, file_location: str) -> int:
    file.file.seek(0)
    total_size = 0
    max_bytes = _max_upload_bytes()

    with open(file_location, "wb") as buffer:
        while True:
            chunk = file.file.read(UPLOAD_CHUNK_SIZE)
            if not chunk:
                break

            total_size += len(chunk)
            if total_size > max_bytes:
                raise HTTPException(
                    status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    detail=f"File is too large. Maximum allowed size is {settings.MAX_UPLOAD_SIZE_MB} MB.",
                )

            buffer.write(chunk)

    return total_size


def _count_csv_records(file_path: str) -> int:
    line_count = 0
    file_size = os.path.getsize(file_path)

    with open(file_path, "rb") as raw_file:
        for chunk in iter(lambda: raw_file.read(UPLOAD_CHUNK_SIZE), b""):
            line_count += chunk.count(b"\n")

        if file_size:
            raw_file.seek(file_size - 1)
            if raw_file.read(1) != b"\n":
                line_count += 1

    return max(line_count - 1, 0)


def _inspect_csv(file_path: str, metrics: str) -> FileInspection:
    with open(file_path, newline="", encoding="utf-8-sig", errors="replace") as csv_file:
        reader = csv.reader(csv_file)
        columns = _normalize_columns(next(reader, []))

        if metrics == "strategy":
            validate_strategy_rows(columns, reader)

    return FileInspection(columns=columns, records_count=_count_csv_records(file_path))


def _inspect_excel(file_path: str, metrics: str) -> FileInspection:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        columns = _normalize_columns(next(rows, []))

        if metrics == "strategy":
            validate_strategy_rows(columns, rows)

        max_row = sheet.max_row or 0
        records_count = max(max_row - 1, 0) if max_row else None
        return FileInspection(columns=columns, records_count=records_count)
    finally:
        workbook.close()


def _inspect_saved_file(file_path: str, file_extension: str, metrics: str) -> FileInspection:
    if file_extension == ".csv":
        inspection = _inspect_csv(file_path, metrics)
    else:
        inspection = _inspect_excel(file_path, metrics)

    if metrics == "analytics":
        validate_analytics_columns(inspection.columns)
    elif metrics == "prediction":
        validate_prediction_columns(inspection.columns)
    elif metrics == "strategy":
        validate_strategy_columns(inspection.columns)

    return inspection


@router.post("/uploadfile/", response_model=FilePublic, status_code=status.HTTP_201_CREATED)
async def upload_file(
    displayname: str = Form(),
    description: str = Form(""),
    file: UploadFile = File(...),
    decisionMetrics: str = Form(),
    db: Session = Depends(get_db),
    current_user: AudienceResponse = Depends(get_current_user)
):
    """
    Uploads a single file to the server and records its metadata in the database.
    - Requires authentication.
    - Stores the file in the configured UPLOAD_DIRECTORY.
    - Returns the metadata of the uploaded file including its ID.
    """
    # Ensure the upload directory exists (use absolute path)
    upload_dir_abs = os.path.abspath(settings.UPLOAD_DIRECTORY)
    os.makedirs(upload_dir_abs, exist_ok=True)

    # Sanitize filename to prevent directory traversal issues
    orig_filename = os.path.basename(file.filename)
    unique_suffix = str(uuid.uuid4())
    name_part, file_extension = os.path.splitext(orig_filename)
    filename = f"{name_part}_{unique_suffix}{file_extension}"

    # ✅ FIX: Store as absolute path so os.path.exists() always resolves correctly
    # regardless of the working directory the backend process was launched from.
    file_location = os.path.abspath(os.path.join(settings.UPLOAD_DIRECTORY, filename))

    metrics = decisionMetrics  # Treat decisionMetrics as a plain string

    try:
        # Restrict allowed file types
        allowed_extensions = [".csv", ".xlsx"]
        if file_extension.lower() not in allowed_extensions:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Only CSV and Excel (.xlsx) files are allowed."
            )

        file_size = await run_in_threadpool(_copy_upload_to_disk, file, file_location)
        inspection = await run_in_threadpool(
            _inspect_saved_file,
            file_location,
            file_extension.lower(),
            metrics,
        )
        file_type = file.content_type or "application/octet-stream"

        # Create a new file record in the database
        db_file = crud_files.create_file_record(
            db=db,
            filename=filename,
            filepath=file_location,   # absolute path stored here
            displayname=displayname,
            description=description,
            file_size=file_size,
            file_type=file_type,
            owner_id=current_user.id,
            decisionMetrics=metrics,
            records=inspection.records_count
        )
        return db_file
    except HTTPException:
        if os.path.exists(file_location):
            os.remove(file_location)
        raise
    except Exception as e:
        # Clean up the partially uploaded file if database record creation fails
        if os.path.exists(file_location):
            os.remove(file_location)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Could not upload file or record metadata: {e}"
        )
@router.get("/files/", response_model=List[FilePublic])
async def list_files(
    decisionMetrics: str = Query(None, description="Filter by decisionMetrics (plain string, e.g. 'analytics', 'prediction', 'strategy')"),
    db: Session = Depends(get_db),
    current_user: AudienceResponse = Depends(get_current_user)
):
    """
    Lists all uploaded files for the current user, fetched from the database.
    - Requires authentication.
    - Returns a list of file metadata (including IDs).
    """
    # Fetch files associated with the current user from the database
    files = crud_files.get_files_by_owner(db, current_user.id)
    # If decisionMetrics is provided, filter files in-memory by plain string match
    if decisionMetrics:
        files = [f for f in files if getattr(f, 'decisionMetrics', None) == decisionMetrics or getattr(f, 'decision_metrics', None) == decisionMetrics]
    return files

@router.delete("/files/")
async def delete_current_user_files(
    db: Session = Depends(get_db),
    current_user: AudienceResponse = Depends(get_current_user)
):
    """
    Deletes all uploaded files for the current user across all features.
    Called during logout so a fresh session starts with no previous uploads.
    """
    files = crud_files.get_files_by_owner(db, current_user.id, limit=None)
    removed_from_disk = 0

    for db_file in files:
        file_path = db_file.path
        if file_path and _is_safe_upload_path(file_path) and os.path.exists(file_path):
            os.remove(file_path)
            removed_from_disk += 1

    deleted_records = crud_files.delete_files_by_owner(db, current_user.id)
    return {
        "deleted_records": len(deleted_records),
        "removed_from_disk": removed_from_disk,
    }

@router.get("/files/{file_id}")
async def serve_file(
    file_id: int, # Now fetching by ID
    db: Session = Depends(get_db), # Inject database session
    current_user: AudienceResponse = Depends(get_current_user) # Protect this endpoint
):
    """
    Serves a specific uploaded file by its ID.
    - Requires authentication.
    - Returns the file as a FileResponse.
    - Ensures the user owns the file before serving.
    """
    # Retrieve file metadata from the database using the file_id
    db_file = crud_files.get_file_by_id(db, file_id)
    if not db_file:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found in database")
    
    # Authorize: Ensure the current user is the owner of the file
    if db_file.audience_id != current_user.id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not authorized to access this file")

    file_path = db_file.path # Use the stored file path

    if not os.path.exists(file_path):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found on disk (metadata exists)")
    
    # Ensure the file is within the allowed upload directory to prevent path traversal
    if not _is_safe_upload_path(file_path):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid file path")

    return FileResponse(path=file_path, filename=db_file.filename, media_type="application/octet-stream")

def validate_analytics_columns(columns: List[str]):
    required_cols = ["timestamp", "endpoint", "client_id", "requests", "status_code"]
    missing = [col for col in required_cols if col not in columns]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Missing required columns for analytics: {', '.join(missing)}"
        )


def validate_prediction_columns(columns: List[str]):
    required_fields = [
        {"name": "timestamp", "mappedKeys": ["date", "datetime", "time", "created_at", "timestamp"]},
        {"name": "request_count", "mappedKeys": ["requests", "request_count", "api_requests", "calls", "hits"]}
    ]
    missing = []
    for field in required_fields:
        if not any(col in columns for col in field["mappedKeys"]):
            missing.append(field["name"])
    if missing:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Missing required columns for prediction: {', '.join(missing)}"
        )


def validate_strategy_columns(columns: List[str]):
    if len(columns) != 2 or set(columns) != {"question", "answer"}:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Strategy template must have exactly two columns: 'question' and 'answer'"
        )


def validate_strategy_rows(columns: List[str], rows: Iterable[Iterable[object]]):
    allowed_answers = ["Yes", "No", "Maybe"]
    allowed_set = set(allowed_answers)
    validate_strategy_columns(columns)
    answer_index = columns.index("answer")
    invalid_answers = []
    seen_invalid = set()

    for row in rows:
        row_values = list(row)
        if not any(value is not None and str(value).strip() for value in row_values):
            continue

        answer = ""
        if answer_index < len(row_values) and row_values[answer_index] is not None:
            answer = str(row_values[answer_index]).strip()

        if answer not in allowed_set and answer not in seen_invalid:
            invalid_answers.append(answer or "<blank>")
            seen_invalid.add(answer)

        if len(invalid_answers) >= 10:
            break

    if invalid_answers:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Invalid answers found in strategy template: {', '.join(invalid_answers)}. Allowed answers are: {', '.join(allowed_answers)}"
        )
